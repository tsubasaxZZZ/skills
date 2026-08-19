#!/usr/bin/env python3
"""Optional export from tsr-investigation.yaml. Do not run unless the user asked.

  python3 export.py xlsx
  python3 export.py csv
  python3 export.py md
"""

from __future__ import annotations

import argparse
import csv
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib import (
    STATUS_LABELS,
    load_config,
    load_investigation,
    project_root,
    python_has_module,
    venv_python,
    xlsx_env_plan,
)

COLUMNS = [
    ("id", "ID"),
    ("section", "Section"),
    ("index", "Index"),
    ("title", "Title"),
    ("priority", "Priority"),
    ("severity", "Severity"),
    ("status", "Status"),
    ("report_claim", "Report claim"),
    ("evidence", "Evidence"),
    ("next_actions", "Next actions"),
    ("investigated_on", "Investigated on"),
    ("user_notes", "User notes"),
    ("decision_brief.scope", "Scope"),
    ("decision_brief.if_ignored", "If ignored"),
    ("decision_brief.remediation_cost", "Remediation cost"),
    ("decision_brief.depends_on", "Depends on"),
    ("decision_brief.relieves", "Relieves"),
    ("decision_brief.decision_changes_if", "Decision changes if"),
    ("decision_brief.unconfirmed", "Unconfirmed"),
    ("decision_brief.user_decision", "Decision"),
]


def _status_label(value: str | None) -> str:
    if not value:
        return STATUS_LABELS["not_started"]
    return STATUS_LABELS.get(value, value)


def _cell(finding: dict, key: str) -> str:
    if key == "status":
        return _status_label(finding.get("status"))
    cur: object = finding
    for part in key.split("."):
        if not isinstance(cur, dict):
            return ""
        cur = cur.get(part)
    if key == "next_actions" or key.endswith("depends_on") or key.endswith("relieves"):
        if isinstance(cur, list):
            return "; ".join(str(a) for a in cur)
        if cur is None:
            return ""
        return str(cur)
    if cur is None:
        return ""
    return str(cur)


def rows_from(data: dict) -> list[dict]:
    return list(data.get("findings") or [])


def export_csv(data: dict, dest: Path) -> None:
    findings = rows_from(data)
    with dest.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow([label for _, label in COLUMNS])
        for f in findings:
            writer.writerow([_cell(f, key) for key, _ in COLUMNS])


def export_md(data: dict, dest: Path) -> None:
    meta = data.get("meta") or {}
    findings = rows_from(data)
    lines = [
        "# TSR investigation",
        "",
        f"- Cluster: {meta.get('cluster')} ({meta.get('cluster_id')})",
        f"- OCP: {meta.get('ocp_version')}",
        f"- Report ID: {meta.get('report_id')}",
        f"- Data quality: {meta.get('data_quality_notice') or ''}",
        "",
        "| ID | Priority | Status | Title |",
        "|----|----------|------------|----------|",
    ]
    for f in findings:
        title = (f.get("title") or "").replace("|", "/")
        lines.append(
            f"| {f.get('id')} | {f.get('priority')} | {_status_label(f.get('status'))} | {title} |"
        )
    env = meta.get("environment") or {}
    if env:
        lines.extend(
            [
                "",
                "## Environment",
                "",
                f"- Role: {env.get('role') or ''}",
                f"- Workload impact: {env.get('workload_impact') or ''}",
                f"- Ops: {env.get('ops') or ''}",
                f"- Upcoming changes: {env.get('upcoming_changes') or ''}",
            ]
        )
    lines.extend(
        [
            "",
            "## Decision brief",
            "",
            "| Finding | Status | Scope | Remediation cost | Depends on | Decision | Decision changes if |",
            "|---------|--------|-------|------------------|------------|----------|---------------------|",
        ]
    )
    for f in findings:
        lines.append(
            "| "
            + " | ".join(
                [
                    str(f.get("id") or ""),
                    _status_label(f.get("status")),
                    _cell(f, "decision_brief.scope").replace("|", "/"),
                    _cell(f, "decision_brief.remediation_cost").replace("|", "/"),
                    _cell(f, "decision_brief.depends_on").replace("|", "/"),
                    _cell(f, "decision_brief.user_decision").replace("|", "/"),
                    _cell(f, "decision_brief.decision_changes_if").replace("|", "/"),
                ]
            )
            + " |"
        )
    dest.write_text("\n".join(lines) + "\n", encoding="utf-8")


def require_openpyxl(root: Path) -> None:
    """Use current interpreter, or the project venv. Never pip install --user."""
    if python_has_module(sys.executable, "openpyxl"):
        return
    vpy = venv_python(root)
    if vpy.is_file() and python_has_module(vpy, "openpyxl"):
        raise SystemExit(subprocess.run([str(vpy), *sys.argv]).returncode)
    plan = xlsx_env_plan(root)
    commands = "\n".join(f"  {c}" for c in plan["commands"])
    raise SystemExit(
        "openpyxl is required for xlsx export. Do not use pip install --user.\n"
        f"Install method: {plan['method']} (uv if available, otherwise venv).\n"
        "Ask the user, then run setup.py ensure-xlsx in this skill's scripts directory.\n"
        f"Equivalent commands:\n{commands}"
    )


def export_xlsx(data: dict, dest: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    findings = rows_from(data)
    meta = data.get("meta") or {}
    wb = Workbook()

    ws = wb.active
    ws.title = "Findings"
    headers = [label for _, label in COLUMNS]
    ws.append(headers)
    for f in findings:
        ws.append([_cell(f, key) for key, _ in COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(findings) + 1, 2)}"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    widths = [28, 32, 8, 56, 12, 10, 16, 60, 40, 32, 14, 32, 40, 40, 40, 24, 24, 40, 32, 20]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    summary = wb.create_sheet("Summary")
    summary.append(["Field", "Value"])
    summary.append(["Cluster", meta.get("cluster") or ""])
    summary.append(["Cluster ID", meta.get("cluster_id") or ""])
    summary.append(["OCP Version", meta.get("ocp_version") or ""])
    summary.append(["Report ID", meta.get("report_id") or ""])
    summary.append(["Data collected", meta.get("data_collected") or ""])
    summary.append(["Must-gather", meta.get("must_gather_label") or ""])
    summary.append(["Data Quality Notice", meta.get("data_quality_notice") or ""])
    env = meta.get("environment") or {}
    if env:
        summary.append([])
        summary.append(["Environment", ""])
        summary.append(["Role", env.get("role") or ""])
        summary.append(["Workload impact", env.get("workload_impact") or ""])
        summary.append(["Ops", env.get("ops") or ""])
        summary.append(["Upcoming changes", env.get("upcoming_changes") or ""])
    summary.append([])
    summary.append(["Priority", "Count"])
    pri_counts: dict[str, int] = {}
    st_counts: dict[str, int] = {}
    for f in findings:
        pri_counts[f.get("priority") or "?"] = pri_counts.get(f.get("priority") or "?", 0) + 1
        label = _status_label(f.get("status"))
        st_counts[label] = st_counts.get(label, 0) + 1
    for k, v in pri_counts.items():
        summary.append([k, v])
    summary.append([])
    summary.append(["Status", "Count"])
    for k, v in st_counts.items():
        summary.append([k, v])
    summary.append([])
    summary.append(["Section", "Count"])
    for cat in data.get("categories") or []:
        n = sum(1 for f in findings if f.get("section") == cat)
        summary.append([cat, n])
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 80

    wb.save(dest)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("format", choices=("xlsx", "csv", "md"))
    parser.add_argument("--root", help="Project root")
    parser.add_argument("--out", help="Output file path")
    args = parser.parse_args()

    root = project_root(args.root)
    cfg = load_config(root) or {}
    paths = cfg.get("paths") or {}
    src = Path(paths.get("investigation") or root / "tsr-investigation.yaml")
    export_dir = Path(paths.get("export_dir") or root)
    data = load_investigation(src)
    if not data:
        raise SystemExit(f"investigation YAML not found: {src}. Run seed.py first.")

    default_name = {
        "xlsx": "tsr-investigation.xlsx",
        "csv": "tsr-investigation.csv",
        "md": "tsr-investigation.md",
    }[args.format]
    dest = Path(args.out).expanduser() if args.out else export_dir / default_name

    if args.format == "csv":
        export_csv(data, dest)
    elif args.format == "md":
        export_md(data, dest)
    else:
        require_openpyxl(root)
        export_xlsx(data, dest)
    print(f"wrote {dest}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
